"""Test helpers: a PDF and a PPTX with real extractable text, and the text
the KNOWLEDGE_EMBED job chunks from an artifact's bytes (review
chunk-data-12: raw-slice readers must decode as that job decoded)."""

from __future__ import annotations

import hashlib
import io
from typing import Any


def pdf_bytes(text: str) -> bytes:
    """A one-page PDF whose page shows ``text`` (Helvetica, one Tj)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    trailer = b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
    return out + trailer % (len(objects) + 1, xref)


def pptx_bytes(text: str) -> bytes:
    """A one-slide deck whose text box holds ``text``."""
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text = text
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


class _S3:
    def __init__(self, raw: bytes):
        self.raw = raw

    def get_object(self, **_bucket_and_key: str) -> dict[str, Any]:
        return {"Body": io.BytesIO(self.raw)}


def embed_text(raw: bytes, source_name: str) -> str:
    """The text KNOWLEDGE_EMBED chunks for an artifact of these bytes."""
    from umrs.jobs import knowledge_adapters as ka

    src: dict[str, Any] = {
        "artifact_id": "art-x",
        "source_ref": "k",
        "kind": "document",
        "metadata": {"source_name": source_name},
    }
    text = ka._artifact_text(src, _S3(raw), "b", ka.KnowledgeEmbedResult())
    assert text, f"the embed job read no text from {source_name}"
    return text


def chunk_content(document: str, start: int, end: int) -> dict[str, Any]:
    """A chunk node's content for ``document[start:end]``."""
    body = document[start:end]
    return {
        "kind": "chunk",
        "artifact_id": "art-x",
        "chunk_index": 0,
        "chunk_start": start,
        "chunk_end": end,
        "chunk_md5": hashlib.md5(body.encode("utf-8")).hexdigest(),
        "excerpt": body[:10],
    }
