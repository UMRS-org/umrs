"""Cycle-3 memfix #2 (2026-10-22): a PDF or PPTX between READ_CAP (64 MiB)
and the 512 MiB upload cap is understood and embedded again.

Found by the mem review (``.stab-cycle3/memreview/memreview-v1-1``): after
cycle-3 P1, ``_read_bounded`` read a binary document whole only within
READ_CAP and refused a larger one by name, so an 80 MiB thesis that
processed before (read whole) failed understand, was dropped from the
chained embed and was recorded as a permanent embed failure, while the
upload still answered 200. Now such a document is streamed to a temporary
file in 1 MiB parts and its text is extracted from there: pypdf reads only
the objects the pages use, and a deck is copied without its media first
(python-pptx loads every part's bytes).
"""

from __future__ import annotations

import io
import os
import tracemalloc
import zlib
from typing import Any

import pytest
from tests.test_c3_knowledge_bounded_reads import _S3
from umrs.data_catalog import extract
from umrs.jobs import knowledge_adapters as ka
from umrs.knowledge import canonical_identity

_TEXT = "Thesis chapter one: quantum materials"
_KEY = ("kb", "projects/p/knowledge/artifacts/h")


def _pdf(pad: int, text: str = _TEXT) -> bytes:
    """A valid one-page PDF showing *text*, plus *pad* bytes of an unused
    (incompressible) stream object — a figure the page never draws."""
    shown = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    filler = os.urandom(pad)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(shown) + shown + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(filler) + filler + b"\nendstream",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.7\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    out.write(b"".join(b"%010d 00000 n \n" % offset for offset in offsets))
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def _pdf_with_figures(pages: int, image_bytes: int) -> bytes:
    """*pages* text pages, each drawing its own raw RGB image."""
    side = int((image_bytes / 3) ** 0.5)
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    kids = []
    for page in range(pages):
        number = 4 + 3 * page
        kids.append(b"%d 0 R" % number)
        objects.update(_figure_page(number, side, f"{_TEXT}, page {page + 1}"))
    objects[2] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(kids), pages)
    return _pdf_file(objects)


def _figure_page(number: int, side: int, text: str) -> dict[int, bytes]:
    shown = f"q {side} 0 0 {side} 0 0 cm /Im0 Do Q BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
    image = os.urandom(side * side * 3)
    return {
        number: b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R"
        b" /Resources << /Font << /F1 3 0 R >> /XObject << /Im0 %d 0 R >> >> >>"
        % (number + 1, number + 2),
        number + 1: _stream(b"", shown.encode()),
        number + 2: _stream(
            b"/Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceRGB"
            b" /BitsPerComponent 8 " % (side, side),
            image,
        ),
    }


def _stream(fields: bytes, data: bytes) -> bytes:
    return b"<< %s/Length %d >>\nstream\n" % (fields, len(data)) + data + b"\nendstream"


def _pdf_file(objects: dict[int, bytes]) -> bytes:
    out = io.BytesIO()
    out.write(b"%PDF-1.7\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = out.tell()
        out.write(b"%d 0 obj\n" % number + objects[number] + b"\nendobj\n")
    xref, size = out.tell(), max(objects) + 1
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % size)
    out.write(b"".join(b"%010d 00000 n \n" % offsets.get(n, 0) for n in range(1, size)))
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref))
    return out.getvalue()


def _png(side: int) -> bytes:
    """An incompressible RGB PNG of side x side pixels."""
    raw = b"".join(b"\x00" + os.urandom(side * 3) for _ in range(side))

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return len(data).to_bytes(4, "big") + body + zlib.crc32(body).to_bytes(4, "big")

    header = side.to_bytes(4, "big") * 2 + b"\x08\x02\x00\x00\x00"
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw, 0))
        + chunk(b"IEND", b"")
    )


def _pptx(image_side: int) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Graphene membranes"
    slide.placeholders[1].text = _TEXT
    slide.shapes.add_picture(io.BytesIO(_png(image_side)), Inches(1), Inches(1))
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def _source(name: str) -> dict[str, Any]:
    return {"artifact_id": "art-1", "source_ref": _KEY[1], "metadata": {"source_name": name}}


@pytest.fixture
def small_cap(monkeypatch):
    monkeypatch.setattr(ka, "_UNDERSTAND_PREFIX_BYTES", 1024)
    monkeypatch.setattr(canonical_identity, "READ_CAP", 4096)


def test_understand_extracts_a_pdf_past_read_cap_from_disk(small_cap):
    s3 = _S3({_KEY: _pdf(64 * 1024)})
    result = ka.KnowledgeUnderstandResult()
    content = ka._understand_bytes(s3, "kb", _source("thesis.pdf"), result)
    assert result.failed_items == []
    assert content is not None and _TEXT in content.decode()
    assert max(s3.reads) <= ka._STREAM_PART_BYTES


def test_embed_reads_the_text_of_a_pdf_past_read_cap(small_cap):
    s3 = _S3({_KEY: _pdf(ka._embed_read_bytes() + 64 * 1024)})
    result = ka.KnowledgeEmbedResult()
    text = ka._artifact_text(_source("thesis.pdf"), s3, "kb", result)
    assert result.failed_items == [] and text is not None and _TEXT in text


def test_a_deck_past_read_cap_gives_the_text_of_the_whole_deck(small_cap):
    deck = _pptx(256)
    s3 = _S3({_KEY: deck})
    result = ka.KnowledgeUnderstandResult()
    content = ka._understand_bytes(s3, "kb", _source("talk.pptx"), result)
    assert result.failed_items == [] and content is not None
    whole, _tags = extract._extract_pptx_detailed(deck)
    assert content.decode() == whole and _TEXT in whole


def test_a_pdf_with_no_text_past_read_cap_fails_by_name(small_cap):
    s3 = _S3({_KEY: _pdf(ka._embed_read_bytes() + 4096, text="")})
    result = ka.KnowledgeEmbedResult()
    assert ka._embed_bytes(_source("scan.pdf"), s3, "kb", result) is None
    (failed,) = result.failed_items
    assert "no extractable text" in failed["error"] and failed["permanent"]


def test_a_pdf_past_the_spool_cap_fails_by_name(small_cap, monkeypatch):
    monkeypatch.setattr(ka, "_DOCUMENT_SPOOL_MAX", 32 * 1024)
    s3 = _S3({_KEY: _pdf(64 * 1024)})
    result = ka.KnowledgeUnderstandResult()
    assert ka._understand_bytes(s3, "kb", _source("thesis.pdf"), result) is None
    assert "too large to extract" in result.failed_items[0]["error"]
    assert sum(s3.reads) <= 1025  # the ranged head only


def _peak_while(call: Any) -> tuple[Any, int]:
    tracemalloc.start()
    try:
        out = call()
        return out, tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()


@pytest.mark.parametrize(
    ("name", "make"),
    [
        ("thesis.pdf", lambda: _pdf(24 << 20)),
        ("figures.pdf", lambda: _pdf_with_figures(8, 3 << 20)),
        ("talk.pptx", lambda: _pptx(2900)),
    ],
)
def test_a_large_document_is_extracted_in_bounded_memory(small_cap, name, make):
    """A 24 MiB document: its bytes never enter memory whole."""
    body = make()
    assert len(body) > 20 << 20
    s3 = _S3({_KEY: body})
    result = ka.KnowledgeUnderstandResult()
    content, peak = _peak_while(lambda: ka._understand_bytes(s3, "kb", _source(name), result))
    assert result.failed_items == [] and _TEXT in content.decode()
    assert peak < 8 << 20, peak
