"""Text and metadata extraction for catalog ingestion (MVP, no GPU)."""

from __future__ import annotations

import csv
import io
import itertools
import json
import os
import shutil
import tempfile
import zipfile
from typing import Any, BinaryIO

from .modality import detect_modality


def extract_content(
    filename: str,
    content: bytes,
) -> tuple[str, str, list[str]]:
    """
    Extract searchable text and a short summary from object bytes.

    Returns:
        (summary, extracted_text, tags)
    """
    modality = detect_modality(filename, content)
    if modality == "text":
        text = _decode_text(content)
        return _summarize(text), text, []
    if modality == "structured":
        text, tags = _extract_structured(filename, content)
        return _summarize(text), text, tags
    if modality == "document":
        text, tags = _extract_document_detailed(filename, content)
        return _summarize(text), text, tags
    if modality == "image":
        text = _extract_image_metadata(content)
        summary = text or f"Image file {filename}"
        return summary[:500], text, []
    if modality == "notebook":
        from ..knowledge.parse_notebook import parse_notebook

        result = parse_notebook(filename, content)
        return result.summary, result.extracted_text, result.tags
    if modality == "code":
        from ..knowledge.parse_code import parse_code

        result = parse_code(filename, content)
        return result.summary, result.extracted_text, result.tags
    return f"{modality} file {filename}", "", []


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


def _extract_structured(filename: str, content: bytes) -> tuple[str, list[str]]:
    text = _decode_text(content)
    tags: list[str] = []
    lower = filename.lower()
    if lower.endswith(".json"):
        try:
            payload = json.loads(text)
            tags.extend(_json_keys(payload, max_keys=12))
            text = json.dumps(payload, indent=2)[:50_000]
        except json.JSONDecodeError:
            pass
    elif lower.endswith((".csv", ".tsv")):
        delimiter = "\t" if lower.endswith(".tsv") else ","
        try:
            reader = csv.reader(io.StringIO(text), delimiter=delimiter)
            rows = list(reader)
            if rows:
                header = rows[0][:12]
                tags.extend(header)
                for col_idx in range(min(3, len(header))):
                    seen: set[str] = set()
                    for row in rows[1:201]:
                        if col_idx >= len(row):
                            continue
                        value = row[col_idx].strip()
                        if not value or value in seen:
                            continue
                        seen.add(value)
                        tags.append(value)
                        if len(seen) >= 5:
                            break
                text = "\n".join(",".join(row) for row in rows[:200])
        except csv.Error:
            pass
    return text, [t for t in tags if t][:24]


def _json_keys(payload: Any, *, max_keys: int, prefix: str = "") -> list[str]:
    keys: list[str] = []
    if isinstance(payload, dict):
        for key, value in payload.items():
            keys.append(f"{prefix}{key}" if prefix else str(key))
            if len(keys) >= max_keys:
                break
            keys.extend(
                _json_keys(value, max_keys=max_keys - len(keys), prefix=f"{key}.")
            )
    elif isinstance(payload, list) and payload and isinstance(payload[0], dict):
        keys.extend(_json_keys(payload[0], max_keys=max_keys))
    return keys[:max_keys]


def _extract_document_detailed(filename: str, content: bytes) -> tuple[str, list[str]]:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _extract_pdf_detailed(content)
    if lower.endswith((".pptx", ".ppt")):
        return _extract_pptx_detailed(content)
    return "", ["extract_unsupported"]


def _extract_pdf_detailed(content: bytes) -> tuple[str, list[str]]:
    return _pdf_text(io.BytesIO(content))


def _pdf_text(stream: BinaryIO) -> tuple[str, list[str]]:
    """A PDF's text from a seekable stream; pypdf reads only the objects
    the pages it extracts use, and they are dropped after each page: its
    object cache kept every image a page draws (a 400 MiB figure-heavy
    PDF read from disk still took +420 MB; cycle-3 memfix #2)."""
    try:
        from pypdf import PdfReader  # type: ignore[import-untyped]
    except ImportError:
        return "", ["extract_dependency_missing"]

    try:
        reader = PdfReader(stream)
        parts = []
        max_pages = _max_document_pages()
        for page in reader.pages[:max_pages]:
            parts.append(page.extract_text() or "")
            _drop_resolved_objects(reader)
        text = "\n".join(parts)[: _max_document_text_chars()]
    except Exception:
        return "", ["extract_failed"]

    if not text.strip():
        return "", ["extract_empty"]
    return text, []


def _drop_resolved_objects(reader: Any) -> None:
    """Empty pypdf's resolved-object cache (re-read on the next use)."""
    cache = getattr(reader, "resolved_objects", None)
    if isinstance(cache, dict):
        cache.clear()


def _extract_pptx_detailed(content: bytes) -> tuple[str, list[str]]:
    return _pptx_text(io.BytesIO(content))


def _pptx_text(stream: BinaryIO) -> tuple[str, list[str]]:
    try:
        from pptx import Presentation  # type: ignore[import-untyped]
    except ImportError:
        return "", ["extract_dependency_missing"]

    try:
        presentation = Presentation(stream)
        parts = []
        max_slides = _max_document_pages()
        # Slides support iteration, not slicing (python-pptx 1.0 raises on
        # [:n], which turned every deck into "extract_failed").
        for slide in itertools.islice(presentation.slides, max_slides):
            for shape in slide.shapes:
                text = getattr(shape, "text", None)
                if text:
                    parts.append(text)
        text = "\n".join(parts)[: _max_document_text_chars()]
    except Exception:
        return "", ["extract_failed"]

    if not text.strip():
        return "", ["extract_empty"]
    return text, []


def extract_document_file(
    kind: str, spool: BinaryIO, *, spool_dir: str | None = None
) -> tuple[str, list[str]]:
    """``(text, tags)`` of a PDF (``kind`` "pdf") or a PPTX held in a
    seekable file, never read whole into memory (cycle-3 memfix #2: a
    document past READ_CAP is streamed to disk, not refused). pypdf reads
    the objects the extracted pages use; python-pptx loads every part's
    bytes, so a deck is first copied (in ``spool_dir``) without the bytes
    of its non-XML parts — media, embeddings, thumbnails, the bulk of a
    deck. Its text is all in the XML parts, so it is the same text.

    A failed write of that copy raises :class:`SpoolWriteError` (an
    ``OSError``): it is the spool disk, not the deck (fixverify memfix-r3
    round 2: a full disk was "extract_failed", a permanent refusal)."""
    if kind == "pdf":
        return _pdf_text(spool)
    with tempfile.TemporaryFile(dir=spool_dir) as slim:
        try:
            _copy_xml_parts(spool, slim)
        except SpoolWriteError:
            raise
        except (zipfile.BadZipFile, OSError, ValueError):
            return "", ["extract_failed"]
        slim.seek(0)
        return _pptx_text(slim)


class SpoolWriteError(OSError):
    """A write to a spool file failed (a full or read-only spool disk)."""


class _SpoolWriter:
    """A spool file whose failed writes raise :class:`SpoolWriteError`."""

    def __init__(self, target: BinaryIO) -> None:
        self._target = target

    def write(self, data: bytes) -> int:
        try:
            return self._target.write(data)
        except OSError as exc:
            raise SpoolWriteError(exc.errno, exc.strerror or str(exc)) from exc

    def __getattr__(self, name: str) -> Any:
        return getattr(self._target, name)


def _copy_xml_parts(source: BinaryIO, target: BinaryIO) -> None:
    """Copy a zip package part by part, streamed; a part that is not XML
    (``.xml`` / ``.rels``) keeps its name with no bytes."""
    writer: Any = _SpoolWriter(target)
    with (
        zipfile.ZipFile(source) as package,
        zipfile.ZipFile(writer, "w", zipfile.ZIP_DEFLATED) as slim,
    ):
        for info in package.infolist():
            if not info.filename.endswith((".xml", ".rels")):
                slim.writestr(info.filename, b"")
                continue
            with package.open(info) as part, slim.open(info.filename, "w") as out:
                shutil.copyfileobj(part, out, 1 << 20)


def _extract_image_metadata(content: bytes) -> str:
    try:
        from PIL import Image
    except ImportError:
        return ""
    try:
        image = Image.open(io.BytesIO(content))
        parts = [f"format={image.format}", f"size={image.size}"]
        exif = image.getexif()
        if exif:
            for tag_id, value in list(exif.items())[:20]:
                parts.append(f"exif:{tag_id}={value}")
        return "; ".join(str(p) for p in parts)
    except Exception:
        return ""


def _summarize(text: str, limit: int = 500) -> str:
    cleaned = " ".join(text.split())
    if not cleaned:
        return ""
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[: limit - 3] + "..."


def _max_document_pages() -> int:
    raw = os.getenv("UMRS_DATA_CATALOG_MAX_DOCUMENT_PAGES", "200")
    try:
        return max(1, int(raw))
    except ValueError:
        return 200


def _max_document_text_chars() -> int:
    raw = os.getenv("UMRS_DATA_CATALOG_MAX_DOCUMENT_TEXT_CHARS", "500000")
    try:
        return max(1_000, int(raw))
    except ValueError:
        return 500_000
